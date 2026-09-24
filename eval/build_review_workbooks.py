"""Build the two Excel files a reviewer (advisor / dietitian) actually opens.

    backend/.venv/Scripts/python.exe eval/build_review_workbooks.py

Writes to eval/review/:

1. ``NutriLift_ตรวจข้อมูล.xlsx`` - every eval question with its reference
   answer, and every knowledge-card section, one row each, with a status
   dropdown and a comment column. Replaces handing over questions.jsonl, which
   a reviewer should never have to read.
2. ``NutriLift_ให้คะแนนคำตอบ_blind.xlsx`` - the blind A/B scoring sheet for
   the 30-question subset the plan (section 4.5) promised, drawn from
   baseline-v10. Which column is RAG stays in the key file only.

Thai text uses Tahoma: it ships with Windows and Office, renders Thai
correctly, and reads as a plain document font (Arial falls back unevenly).
"""

from __future__ import annotations

import csv
import json
import random
import sys
from pathlib import Path

import frontmatter
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from ingest.__main__ import split_sections  # noqa: E402

OUT = ROOT / "eval" / "review"
RUN_ID = "baseline-v10"
SUBSET_SIZE = 30
SEED = 20260924

FONT = "Tahoma"
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="1F2937")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")  # the cells the reviewer fills
FILL_ALT = PatternFill("solid", fgColor="F3F4F6")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

CATEGORY_TH = {
    "protein": "โปรตีน", "energy": "พลังงาน / ลด-เพิ่มน้ำหนัก", "timing": "ช่วงเวลาการกิน",
    "supplement": "อาหารเสริม", "myth": "ความเชื่อผิด ๆ", "thai-food": "อาหารไทย",
    "safety": "ความปลอดภัย", "out-of-scope": "นอกขอบเขต (ต้องปฏิเสธ)",
}
STATUS = ["ถูกต้อง", "ต้องแก้ไข", "ไม่แน่ใจ"]


def load_cards() -> dict[str, frontmatter.Post]:
    cards = {}
    for path in sorted((ROOT / "knowledge" / "cards").glob("*.md")):
        if path.name.startswith("_"):
            continue
        post = frontmatter.load(path)
        cards[post.get("slug") or path.stem] = post
    return cards


def header(ws, row: int, titles: list[str], widths: list[int]) -> None:
    for col, (title, width) in enumerate(zip(titles, widths, strict=True), start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEAD, CENTER, BOX
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(titles))}{row}"


def put(ws, row: int, col: int, value, *, font=F_BODY, align=WRAP_TOP, fill=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font, cell.alignment, cell.border = font, align, BOX
    if fill is not None:
        cell.fill = fill
    return cell


def estimated_height(texts: list[tuple[str, int]]) -> float:
    """Row height so wrapped Thai text is visible without dragging rows open."""
    lines = 1
    for text, width in texts:
        per_line = max(width * 1.1, 10)
        lines = max(lines, sum(max(1, -(-len(p) // int(per_line))) for p in str(text).split("\n")))
    return min(15 * lines + 4, 409)


def instructions(ws, title: str, lines: list[str]) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    ws["B2"] = title
    ws["B2"].font = F_TITLE
    for i, line in enumerate(lines, start=4):
        cell = ws.cell(row=i, column=2, value=line)
        cell.font = F_BOLD if line.startswith("■") else F_BODY
        cell.alignment = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------------------
# 1. data review workbook
# ---------------------------------------------------------------------------


def build_data_review(cards: dict[str, frontmatter.Post]) -> Path:
    questions = [json.loads(line) for line in (ROOT / "eval" / "questions.jsonl").open(encoding="utf-8") if line.strip()]
    wb = Workbook()

    ws = wb.active
    ws.title = "คำชี้แจง"
    instructions(ws, "NutriLift — แบบตรวจสอบความถูกต้องของข้อมูล", [
        "โครงงาน: การพัฒนาแชตบอทให้คำแนะนำด้านโภชนาการเฉพาะบุคคลสำหรับผู้ฝึกเวทเทรนนิ่ง"
        "ด้วยเทคนิคการสร้างคำตอบเสริมด้วยการค้นคืน (RAG)",
        "",
        "■ ไฟล์นี้มี 2 ชีตที่ขอให้ตรวจ",
        f"1) ชีต \"คำถามและเฉลย\" — คำถามทดสอบ {len(questions)} ข้อ พร้อมคำตอบอ้างอิงที่ใช้ให้คะแนนระบบ "
        "ขอให้ตรวจว่าคำตอบอ้างอิงถูกต้องตามหลักวิชาการ",
        f"2) ชีต \"การ์ดความรู้\" — ฐานความรู้ที่แชตบอทใช้ตอบ {len(cards)} การ์ด แยกเป็นหัวข้อย่อยทีละแถว "
        "ขอให้ตรวจว่าเนื้อหาถูกต้องและไม่ชี้นำผิด",
        "",
        "■ วิธีกรอก (กรอกเฉพาะช่องสีเหลือง)",
        "• คอลัมน์ \"ผลตรวจ\" — กดที่ช่องแล้วเลือกจากรายการ: ถูกต้อง / ต้องแก้ไข / ไม่แน่ใจ",
        "• คอลัมน์ \"ความเห็น\" — ถ้าเลือก \"ต้องแก้ไข\" ขอให้ระบุสั้น ๆ ว่าควรแก้เป็นอะไร",
        "• ข้อที่ถูกต้องไม่ต้องเขียนความเห็น",
        "• กรองดูทีละหมวดได้ด้วยปุ่มลูกศรที่หัวตาราง",
        "",
        "■ ตัวอย่างการกรอก",
        "ผลตรวจ: ต้องแก้ไข    ความเห็น: ควรระบุว่าช่วง 1.6-2.2 g/kg ใช้กับผู้ที่ฝึกต่อเนื่อง ไม่ใช่ผู้เริ่มต้นทุกคน",
        "",
        "■ ชีต \"สรุป\" นับจำนวนที่ตรวจแล้วให้อัตโนมัติ",
        "",
        "■ หมายเหตุ",
        "• คำถามหมวด \"นอกขอบเขต\" และ \"ความปลอดภัย\" ระบบต้องปฏิเสธหรือส่งต่อแพทย์ คำตอบอ้างอิงจึงเป็นพฤติกรรมที่คาดหวัง ไม่ใช่ความรู้",
        "• คำถามหมวด \"อาหารไทย\" ที่ถามตัวเลขพลังงาน ระบบดึงจากตารางอาหาร (กรมอนามัย / INMU / ASEAN FCD) ไม่ได้ใช้การ์ด "
        "ช่อง \"การ์ดที่ใช้ตอบ\" จึงว่าง",
        "• ทุกการ์ดมีแหล่งอ้างอิงในคอลัมน์ \"แหล่งอ้างอิงของการ์ด\" ตรวจย้อนได้",
    ])

    qa = wb.create_sheet("คำถามและเฉลย")
    cols = ["รหัส", "หมวด", "คำถาม", "คำตอบอ้างอิง (เฉลย)", "การ์ดที่ใช้ตอบ", "ผลตรวจ", "ความเห็น / สิ่งที่ควรแก้"]
    widths = [8, 16, 38, 70, 26, 12, 40]
    header(qa, 1, cols, widths)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "เลือกจากรายการเท่านั้น", "ค่าไม่ถูกต้อง"
    qa.add_data_validation(dv)
    for i, q in enumerate(questions, start=2):
        titles = "\n".join(str(cards[s].get("title") or s) if s in cards else s for s in q.get("relevant_doc_slugs") or [])
        row_fill = FILL_ALT if i % 2 else None
        values = [q["id"], CATEGORY_TH.get(q["category"], q["category"]), q["question"], q["reference_answer"], titles]
        for col, value in enumerate(values, start=1):
            put(qa, i, col, value, fill=row_fill, align=CENTER if col == 1 else WRAP_TOP)
        put(qa, i, 6, None, fill=FILL_INPUT, align=CENTER)
        put(qa, i, 7, None, fill=FILL_INPUT)
        dv.add(f"F{i}")
        qa.row_dimensions[i].height = estimated_height([(q["question"], 38), (q["reference_answer"], 70), (titles, 26)])
    qa_last = len(questions) + 1

    cd = wb.create_sheet("การ์ดความรู้")
    cols = ["การ์ด", "หัวข้อย่อย", "เนื้อหา", "แหล่งอ้างอิงของการ์ด", "ผลตรวจ", "ความเห็น / สิ่งที่ควรแก้"]
    widths = [22, 22, 80, 40, 12, 36]
    header(cd, 1, cols, widths)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    cd.add_data_validation(dv2)
    row = 2
    for n, (slug, post) in enumerate(cards.items()):
        title = str(post.get("title") or slug)
        sources = "\n".join(f"• {s}" for s in (post.get("sources") or []))
        fill = FILL_ALT if n % 2 else None
        for k, (heading, text) in enumerate(split_sections(post.content)):
            put(cd, row, 1, title if k == 0 else f"({title})", font=F_BOLD if k == 0 else F_BODY, fill=fill)
            put(cd, row, 2, heading or "บทนำ", fill=fill)
            put(cd, row, 3, text, fill=fill)
            put(cd, row, 4, sources if k == 0 else "↑ เหมือนแถวแรกของการ์ด", fill=fill)
            put(cd, row, 5, None, fill=FILL_INPUT, align=CENTER)
            put(cd, row, 6, None, fill=FILL_INPUT)
            dv2.add(f"E{row}")
            cd.row_dimensions[row].height = estimated_height([(text, 80), (sources if k == 0 else "", 40)])
            row += 1
    cd_last = row - 1

    sm = wb.create_sheet("สรุป")
    sm.sheet_view.showGridLines = False
    for col, width in zip("ABCDEF", [3, 28, 14, 14, 14, 14], strict=True):
        sm.column_dimensions[col].width = width
    sm["B2"], sm["B2"].font = "สรุปผลการตรวจ (นับอัตโนมัติ)", F_TITLE
    for col, title in enumerate(["ชีต", "ทั้งหมด", *STATUS, "ยังไม่ตรวจ"], start=2):
        c = sm.cell(row=4, column=col, value=title)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BOX
    sm.column_dimensions["G"].width = 14
    for r, (name, col, last) in enumerate([("คำถามและเฉลย", "F", qa_last), ("การ์ดความรู้", "E", cd_last)], start=5):
        rng = f"'{name}'!{col}2:{col}{last}"
        put(sm, r, 2, name, font=F_BOLD)
        put(sm, r, 3, last - 1, align=CENTER)
        sm.cell(row=r, column=3).comment = Comment("จำนวนแถวในชีตนั้น ใส่ค่าตอนสร้างไฟล์", "NutriLift")
        for k, status in enumerate(STATUS):
            put(sm, r, 4 + k, f'=COUNTIF({rng},"{status}")', align=CENTER)
        put(sm, r, 7, f"=C{r}-SUM(D{r}:F{r})", align=CENTER)

    wb.calculation.fullCalcOnLoad = True
    path = OUT / "NutriLift_ตรวจข้อมูล.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# 2. blind scoring workbook
# ---------------------------------------------------------------------------


def stratified_subset(rows: list[dict], questions: dict[str, dict]) -> list[dict]:
    """SUBSET_SIZE rows, proportional to category, fixed seed, every category kept."""
    by_cat: dict[str, list[dict]] = {}
    for r in rows:
        by_cat.setdefault(questions[r["id"]]["category"], []).append(r)
    total = len(rows)
    quota = {c: max(1, round(SUBSET_SIZE * len(v) / total)) for c, v in by_cat.items()}
    while sum(quota.values()) > SUBSET_SIZE:
        biggest = max(quota, key=lambda c: quota[c])
        quota[biggest] -= 1
    while sum(quota.values()) < SUBSET_SIZE:
        roomy = max(by_cat, key=lambda c: len(by_cat[c]) - quota[c])
        quota[roomy] += 1
    rng = random.Random(SEED)
    picked = []
    for cat in sorted(by_cat):
        picked += rng.sample(by_cat[cat], quota[cat])
    return sorted(picked, key=lambda r: r["id"])


def build_blind() -> tuple[Path, list[str]]:
    questions = {json.loads(line)["id"]: json.loads(line) for line in (ROOT / "eval" / "questions.jsonl").open(encoding="utf-8") if line.strip()}
    rows = list(csv.DictReader((ROOT / "eval" / "reports" / f"{RUN_ID}_expert.csv").open(encoding="utf-8-sig")))
    subset = stratified_subset(rows, questions)

    wb = Workbook()
    ws = wb.active
    ws.title = "คำชี้แจง"
    instructions(ws, "NutriLift — แบบให้คะแนนคำตอบแชตบอท (blind)", [
        f"ขอให้ให้คะแนนคำตอบของแชตบอท {len(subset)} คำถาม คำถามละ 2 คำตอบ (A และ B)",
        "คำตอบ A และ B มาจากระบบ 2 แบบที่ต่างกัน ผู้ให้คะแนนไม่ทราบว่าแบบไหนเป็นแบบไหน เพื่อไม่ให้มีอคติ "
        "(ลำดับ A/B สุ่มแยกทุกข้อ)",
        "",
        "■ เกณฑ์คะแนน 1–5 (กรอกเฉพาะช่องสีเหลือง เลือกจากรายการ)",
        "ความถูกต้อง: 5 = ถูกต้องทั้งหมด · 4 = ถูกเกือบทั้งหมด มีจุดเล็กน้อย · 3 = ถูกบางส่วน · "
        "2 = ผิดเป็นส่วนใหญ่ · 1 = ผิดหรือเป็นอันตราย",
        "ความเป็นประโยชน์: 5 = นำไปใช้ได้ทันที · 4 = มีประโยชน์ · 3 = พอใช้ · 2 = ประโยชน์น้อย · 1 = ไม่มีประโยชน์",
        "",
        "■ ตัวอย่างการกรอก",
        "A ถูกต้อง 5 · B ถูกต้อง 4 · A ประโยชน์ 4 · B ประโยชน์ 5 · ความเห็น: B ตัวเลขถูกแต่ไม่บอกว่าใช้กับผู้หญิงตั้งครรภ์ไม่ได้",
        "",
        "■ หมายเหตุ",
        "• คำถามบางข้อตั้งใจให้อยู่นอกขอบเขตหรือเกี่ยวกับยา/โรค คำตอบที่ดีคือปฏิเสธอย่างสุภาพและแนะนำให้พบผู้เชี่ยวชาญ",
        "• เครื่องหมาย [S1] [S2] ในคำตอบคือเลขอ้างอิงแหล่งข้อมูลของระบบ ไม่ต้องให้คะแนนตัวเลขอ้างอิงเอง",
        "• คอลัมน์ \"คำตอบอ้างอิง\" คือเฉลยที่ผู้วิจัยเตรียมไว้ ใช้ประกอบการพิจารณาได้",
        f"• ชุดนี้สุ่มแบบแบ่งชั้นตามหมวดคำถาม {SUBSET_SIZE} ข้อจากทั้งหมด {len(rows)} ข้อ (seed {SEED})",
    ])

    sc = wb.create_sheet("ให้คะแนน")
    cols = ["รหัส", "คำถาม", "คำตอบอ้างอิง (เฉลย)", "คำตอบ A", "คำตอบ B",
            "A ถูกต้อง (1-5)", "B ถูกต้อง (1-5)", "A ประโยชน์ (1-5)", "B ประโยชน์ (1-5)", "ความเห็น"]
    widths = [7, 26, 34, 55, 55, 10, 10, 10, 10, 30]
    header(sc, 1, cols, widths)
    dlist = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    dlist.error, dlist.errorTitle = "ใส่เลข 1 ถึง 5", "คะแนนไม่ถูกต้อง"
    sc.add_data_validation(dlist)
    for i, r in enumerate(subset, start=2):
        q = questions[r["id"]]
        for col, value in enumerate([r["id"], r["question"], q["reference_answer"], r["answer_A"], r["answer_B"]], start=1):
            put(sc, i, col, value, align=CENTER if col == 1 else WRAP_TOP)
        for col in range(6, 10):
            put(sc, i, col, None, fill=FILL_INPUT, align=CENTER)
            dlist.add(f"{get_column_letter(col)}{i}")
        put(sc, i, 10, None, fill=FILL_INPUT)
        sc.row_dimensions[i].height = estimated_height([(r["answer_A"], 55), (r["answer_B"], 55), (q["reference_answer"], 34)])
    last = len(subset) + 1

    sm = wb.create_sheet("ความคืบหน้า")
    sm.sheet_view.showGridLines = False
    sm.column_dimensions["B"].width = 34
    sm.column_dimensions["C"].width = 14
    sm["B2"], sm["B2"].font = "ความคืบหน้า (นับอัตโนมัติ)", F_TITLE
    put(sm, 4, 2, "คำถามทั้งหมด", font=F_BOLD)
    put(sm, 4, 3, len(subset), align=CENTER)
    put(sm, 5, 2, "ให้คะแนนครบ 4 ช่องแล้ว", font=F_BOLD)
    s = "'ให้คะแนน'!"
    put(sm, 5, 3, f'=SUMPRODUCT(({s}F2:F{last}<>"")*({s}G2:G{last}<>"")*'
                   f'({s}H2:H{last}<>"")*({s}I2:I{last}<>""))', align=CENTER)
    put(sm, 6, 2, "เหลือ", font=F_BOLD)
    put(sm, 6, 3, "=C4-C5", align=CENTER)

    wb.calculation.fullCalcOnLoad = True
    path = OUT / "NutriLift_ให้คะแนนคำตอบ_blind.xlsx"
    wb.save(path)
    return path, [r["id"] for r in subset]


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    cards = load_cards()
    data_path = build_data_review(cards)
    blind_path, ids = build_blind()
    (OUT / "blind_subset_ids.txt").write_text("\n".join(ids) + "\n", encoding="utf-8")
    print("wrote", data_path)
    print("wrote", blind_path)
    print("blind subset:", ", ".join(ids))
    return 0


if __name__ == "__main__":
    sys.exit(main())
