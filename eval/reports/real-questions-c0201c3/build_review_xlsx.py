"""Build an easy-to-review Excel workbook from eval/reports/real-questions-4fca71e/review-th.csv."""
import csv
import math
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
rows = list(csv.DictReader(SRC.open(encoding="utf-8-sig")))
# Optional before-fix run, matched by result id: build_review_xlsx.py new.csv out.xlsx [baseline.csv run-name new-run-name]
BASE = {r["ผลลัพธ์"]: r for r in csv.DictReader(Path(sys.argv[3]).open(encoding="utf-8-sig"))} if len(sys.argv) > 3 else {}
BASE_NAME = sys.argv[4] if len(sys.argv) > 4 else ""
RUN_NAME = sys.argv[5] if len(sys.argv) > 5 else "main 4fca71e (23 ต.ค. 2569) ก่อนการแก้ #32–#46"
# Optional scope/system classification of failed results: FAILURE_TYPES=path/to/failure-types.json
import json, os  # noqa: E401
KIND = json.loads(Path(os.environ["FAILURE_TYPES"]).read_text(encoding="utf8"))["types"] if os.environ.get("FAILURE_TYPES") else {}
KIND_LABEL = {"scope": ("ขอบเขตงาน", "DDEBF7"), "system": ("ข้อบกพร่องระบบ", "F8CBAD")}

FONT = "Tahoma"
f = lambda **k: Font(name=FONT, size=10, **k)
fill = lambda c: PatternFill("solid", start_color=c)
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TOP = Alignment(wrap_text=True, vertical="top")
INPUT = fill("FFF2CC")

AI_LABEL = {
    "pass": ("ผ่าน", "C6EFCE"),
    "pass_by_policy": ("ผ่านตามนโยบาย", "E2EFDA"),
    "fail": ("ไม่ผ่าน", "FFC7CE"),
    "service_failure": ("บริการขัดข้อง", "D9D9D9"),
}
FAIL_TYPE = {
    "": "",
    "not_answered": "ไม่ได้ตอบคำถาม",
    "missing_key_point": "ขาดประเด็นสำคัญ",
    "misread_clarify": "ตีความผิดแล้วถามกลับ",
    "validation_block": "ตัวตรวจปัดตกทั้งคำตอบ",
    "food_lookup_misroute": "ส่งไปค้นฐานข้อมูลอาหารผิด",
    "factual_error": "ข้อมูลผิดหรือขัดแย้ง",
    "over_refusal": "ปฏิเสธเกินจำเป็น",
    "profile_redirect": "ส่งไปกรอกโปรไฟล์แทนการตอบ",
}
CHOICES = ["รอตรวจ", "ผ่าน", "ผ่านบางส่วน", "ไม่ผ่าน"]

# (header, csv key or callable, width, is_input)
COLS = [
    ("รหัส", "ผลลัพธ์", 10, False),
    ("รอบ", "รอบ", 5, False),
    ("คำถามต้นฉบับ", "คำถามต้นฉบับ", 34, False),
    ("พฤติกรรมที่คาดหวัง (ล็อกก่อนรัน)", "พฤติกรรมที่คาดหวัง (ล็อกก่อนรัน)", 34, False),
    ("ห้าม", "ห้าม", 20, False),
    ("คำตอบจริงของแชตบอต", "คำตอบจริง", 95, False),
    ("ผล AI", None, 12, False),
    ("ประเภทที่ไม่ผ่าน", None, 16, False),
    *([("ประเภท: ขอบเขตงาน / ข้อบกพร่องระบบ", None, 16, False), ("เหตุผลของประเภท", None, 30, False)] if KIND else []),
    ("เหตุผล AI", "เหตุผล AI", 34, False),
    *([("ผล AI รอบก่อนแก้", None, 12, False), ("ประเภทที่ไม่ผ่าน รอบก่อนแก้", None, 16, False)] if BASE else []),
    ("ผลตรวจผู้จัดทำ", "ผลตรวจผู้จัดทำ", 13, True),
    ("ผลตรวจผู้เชี่ยวชาญ", "ผลตรวจผู้เชี่ยวชาญ", 13, True),
    ("หมายเหตุ", "หมายเหตุ", 30, True),
    ("แหล่งคำถาม", "แพลตฟอร์ม", 13, False),
    ("การ์ดความรู้ที่อ้าง", "การ์ดความรู้ที่อ้าง", 24, False),
    ("ชุดแกนกลาง", "อยู่ในชุดแกนกลาง", 9, False),
    ("เรียกโมเดล (ครั้ง)", "จำนวนการเรียกโมเดล", 9, False),
    ("เวลา (วินาที)", "เวลา (วินาที ไม่รวมพัก)", 9, False),
]

LETTER = {head: get_column_letter(i) for i, (head, *_rest) in enumerate(COLS, start=1)}
AI, OWNER, EXPERT = LETTER["ผล AI"], LETTER["ผลตรวจผู้จัดทำ"], LETTER["ผลตรวจผู้เชี่ยวชาญ"]


def visual_len(text):
    """Characters that take horizontal space (Thai vowel/tone marks do not)."""
    return sum(1 for ch in text if unicodedata.category(ch) != "Mn")


def lines_needed(text, width):
    per_line = max(1, int(width * 1.05))
    return sum(max(1, math.ceil(visual_len(p) / per_line)) for p in str(text).split("\n"))


wb = Workbook()

# --- Sheet 1: how to review -------------------------------------------------
guide = wb.active
guide.title = "วิธีตรวจ"
guide.column_dimensions["A"].width = 26
guide.column_dimensions["B"].width = 90
lines = [
    ("ตารางตรวจคำตอบจากคำถามจริง 61 ข้อ", None),
    ("ที่มา", f"{SRC.name} — คำตอบของ {RUN_NAME}"),
    *([("รอบก่อนแก้", f"คอลัมน์ \"ผล AI รอบก่อนแก้\" คือผล AI ของคำถามและรอบเดียวกันใน {BASE_NAME} ใช้เทียบเท่านั้น ไม่ต้องตรวจ")] if BASE else []),
    ("จำนวนแถว", f"{len(rows)} ผล (61 คำถาม รอบที่ 1 + ชุดแกนกลาง 20 ข้อ รอบที่ 2 และ 3)"),
    ("", ""),
    ("ช่องที่ต้องกรอก", "เฉพาะ 3 คอลัมน์พื้นเหลืองในชีต \"ตรวจ\": ผลตรวจผู้จัดทำ, ผลตรวจผู้เชี่ยวชาญ, หมายเหตุ "
                         "คอลัมน์อื่นเป็นข้อมูลจากการรัน ไม่ต้องแก้"),
    ("ค่าที่เลือกได้", "รอตรวจ / ผ่าน / ผ่านบางส่วน / ไม่ผ่าน (เลือกจากลูกศรในเซลล์)"),
    ("เกณฑ์", "เทียบคำตอบจริงกับ \"พฤติกรรมที่คาดหวัง\" และ \"ห้าม\" ซึ่งล็อกไว้ก่อนรัน "
              "ผล AI เป็นผลเบื้องต้นเท่านั้น ผู้ตรวจตัดสินเองได้ ถ้าไม่ตรงกับผล AI ให้เขียนเหตุผลในหมายเหตุ"),
    ("ตัวอย่างการกรอก", "ผลตรวจผู้จัดทำ = ผ่านบางส่วน | หมายเหตุ = ตอบเรื่องพลังงานถูก แต่ไม่ได้บอกปริมาณโปรตีนที่ควรได้"),
    ("", ""),
    ("ความหมายของผล AI", ""),
    ("ผ่าน", "ตรงกับพฤติกรรมที่คาดหวัง และไม่มีสิ่งที่ห้าม"),
    ("ผ่านตามนโยบาย", "ปฏิเสธเรื่องการแพทย์ตามนโยบายของผู้จัดทำ ถือว่าถูกต้อง"),
    ("ไม่ผ่าน", "ดูประเภทในคอลัมน์ \"ประเภทที่ไม่ผ่าน\""),
    ("บริการขัดข้อง", "Gemini ตอบผิดพลาดหรือหมดเวลา ไม่ใช่ความผิดของคำตอบ"),
    ("", ""),
    ("ประเภทที่ไม่ผ่าน", ""),
    *[(v, k) for k, v in FAIL_TYPE.items() if k],
    ("", ""),
    ("เคล็ดลับ", "ใช้ปุ่มกรองที่หัวตาราง เช่น กรอง \"ผล AI\" = ไม่ผ่าน เพื่อตรวจข้อที่มีปัญหาก่อน "
                 "คำตอบที่ยาวมากอาจล้นความสูงของแถว คลิกเซลล์แล้วอ่านในแถบสูตร หรือดับเบิลคลิกเซลล์"),
    ("", ""),
    ("สรุปผลตรวจ", "ตารางด้านล่างนับให้อัตโนมัติจากชีต \"ตรวจ\""),
]
for r, (a, b) in enumerate(lines, start=1):
    guide.cell(r, 1, a).font = f(bold=True)
    guide.cell(r, 2, b).font = f()
    guide.cell(r, 2).alignment = TOP
    guide.cell(r, 1).alignment = TOP
guide["A1"].font = Font(name=FONT, size=14, bold=True)
for r, (a, _) in enumerate(lines, start=1):
    if a in ("ช่องที่ต้องกรอก",):
        guide.cell(r, 2).fill = INPUT

n = len(rows) + 1
start = len(lines) + 1
heads = ["ผล", "ผล AI", "ผู้จัดทำ", "ผู้เชี่ยวชาญ"] + (["ผล AI รอบก่อนแก้"] if BASE else [])
for c, h in enumerate(heads, start=1):
    cell = guide.cell(start, c if c < 3 else c, h)
guide.column_dimensions["C"].width = 14
guide.column_dimensions["D"].width = 14
guide.column_dimensions["E"].width = 16
for c, h in enumerate(heads, start=1):
    cell = guide.cell(start, c, h)
    cell.font = f(bold=True, color="FFFFFF")
    cell.fill = fill("1F3864")
    cell.border = BOX
summary = [("ผ่าน", "ผ่าน"), ("ผ่านตามนโยบาย", None), ("ผ่านบางส่วน", "ผ่านบางส่วน"),
           ("ไม่ผ่าน", "ไม่ผ่าน"), ("บริการขัดข้อง", None), ("รอตรวจ", "รอตรวจ")]
for i, (name, human) in enumerate(summary, start=start + 1):
    guide.cell(i, 1, name).font = f()
    guide.cell(i, 2, f'=COUNTIF(ตรวจ!${AI}$2:${AI}${n},"{name}")')
    if BASE:
        b = LETTER["ผล AI รอบก่อนแก้"]
        guide.cell(i, 5, f'=COUNTIF(ตรวจ!${b}$2:${b}${n},"{name}")')
    for col, letter in ((3, OWNER), (4, EXPERT)):
        guide.cell(i, col, f'=COUNTIF(ตรวจ!${letter}$2:${letter}${n},"{human}")' if human else "-")
    for col in range(1, len(heads) + 1):
        guide.cell(i, col).border = BOX
        guide.cell(i, col).font = f()
total = start + len(summary) + 1
guide.cell(total, 1, "รวม").font = f(bold=True)
for col, letter in [(c, get_column_letter(c)) for c in range(2, len(heads) + 1)]:
    guide.cell(total, col, f"=SUM({letter}{start + 1}:{letter}{total - 1})").font = f(bold=True)
for col in range(1, len(heads) + 1):
    guide.cell(total, col).border = BOX
if KIND:
    k = LETTER["ประเภท: ขอบเขตงาน / ข้อบกพร่องระบบ"]
    row0 = total + 2
    guide.cell(row0, 1, "ผลที่ไม่ผ่าน แยกตามประเภท (AI แบ่ง รอผู้ตรวจยืนยัน)").font = f(bold=True)
    for i, (label, _) in enumerate(KIND_LABEL.values(), start=row0 + 1):
        guide.cell(i, 1, label).font = f()
        guide.cell(i, 2, f'=COUNTIF(ตรวจ!${k}$2:${k}${n},"{label}")').font = f()
        for col in (1, 2):
            guide.cell(i, col).border = BOX
    guide.cell(row0 + 1, 3, "คลังความรู้หรือฐานข้อมูลไม่ครอบคลุม ระบบจึงตอบเลี่ยงแทนการแต่ง").font = f()
    guide.cell(row0 + 2, 3, "ระบบตอบผิด ไม่ตรง หรือไม่ครบ ในเรื่องที่มีข้อมูลรองรับ").font = f()

# --- Sheet 2: review grid ---------------------------------------------------
ws = wb.create_sheet("ตรวจ")
for c, (head, _, width, is_input) in enumerate(COLS, start=1):
    cell = ws.cell(1, c, head)
    cell.font = f(bold=True, color="000000" if is_input else "FFFFFF")
    cell.fill = fill("FFD966") if is_input else fill("1F3864")
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = BOX
    ws.column_dimensions[cell.column_letter].width = width
ws.row_dimensions[1].height = 32

for r, row in enumerate(rows, start=2):
    ai, ai_color = AI_LABEL[row["ผล AI เบื้องต้น"]]
    tallest = 1
    for c, (head, key, width, is_input) in enumerate(COLS, start=1):
        base = BASE.get(row["ผลลัพธ์"])
        if head == "ผล AI":
            value = ai
        elif head == "ประเภทที่ไม่ผ่าน":
            value = FAIL_TYPE[row["ประเภทที่ไม่ผ่าน"]]
        elif head == "ประเภท: ขอบเขตงาน / ข้อบกพร่องระบบ":
            kind = KIND.get(row["ผลลัพธ์"])
            value = KIND_LABEL[kind[0]][0] if kind else ""
        elif head == "เหตุผลของประเภท":
            kind = KIND.get(row["ผลลัพธ์"])
            value = kind[1] if kind else ""
        elif head == "ผล AI รอบก่อนแก้":
            value = AI_LABEL[base["ผล AI เบื้องต้น"]][0] if base else "ไม่มี"
        elif head == "ประเภทที่ไม่ผ่าน รอบก่อนแก้":
            value = FAIL_TYPE[base["ประเภทที่ไม่ผ่าน"]] if base else ""
        else:
            value = row[key]
        if head in ("รอบ", "เรียกโมเดล (ครั้ง)"):
            value = int(value)
        elif head == "เวลา (วินาที)":
            value = round(float(value), 1)
        cell = ws.cell(r, c, value)
        cell.font = f()
        cell.alignment = TOP
        cell.border = BOX
        if is_input:
            cell.fill = INPUT
        if head == "ผล AI":
            cell.fill = fill(ai_color)
            cell.font = f(bold=True)
        if head == "ประเภท: ขอบเขตงาน / ข้อบกพร่องระบบ" and KIND.get(row["ผลลัพธ์"]):
            cell.fill = fill(KIND_LABEL[KIND[row["ผลลัพธ์"]][0]][1])
            cell.font = f(bold=True)
        if head == "ผล AI รอบก่อนแก้" and base:
            cell.fill = fill(AI_LABEL[base["ผล AI เบื้องต้น"]][1])
        if head == "แหล่งคำถาม" and row["URL ต้นทาง"]:
            cell.hyperlink = row["URL ต้นทาง"]
            cell.font = f(color="0563C1", underline="single")
        if isinstance(value, str):
            tallest = max(tallest, lines_needed(value, width))
    ws.row_dimensions[r].height = min(409, tallest * 13 + 6)

dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True)
dv.add(f"{OWNER}2:{EXPERT}{n}")
ws.add_data_validation(dv)
for word, color in (("ผ่าน", "C6EFCE"), ("ผ่านบางส่วน", "FFEB9C"), ("ไม่ผ่าน", "FFC7CE")):
    ws.conditional_formatting.add(
        f"{OWNER}2:{EXPERT}{n}", FormulaRule(formula=[f'{OWNER}2="{word}"'], fill=fill(color), stopIfTrue=True))

ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLS)).column_letter}{n}"
ws.sheet_view.zoomScale = 90
wb.active = 1
wb.calculation.fullCalcOnLoad = True
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT, len(rows))
