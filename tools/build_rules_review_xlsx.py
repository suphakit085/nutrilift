"""Build the expert-facing workbook from docs/calculator-rules-review.csv.

The CSV stays the source of truth (a test pins it to the calculator); this
workbook is only a readable form of it: wrapped text, one row per rule, a
drop-down for the verdict and a sign-off sheet. Run from the repo root:

    python tools/build_rules_review_xlsx.py
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "calculator-rules-review.csv"
TARGET = ROOT / "docs" / "calculator-rules-review.xlsx"

#: Columns shown to the reviewer, with widths; the code columns go last.
WIDTHS = {
    "รหัส": 7, "หมวด": 14, "กติกา": 24, "ค่าที่ใช้": 34, "ใช้เมื่อ": 26, "ที่มา": 48, "ประเภท": 18,
    "ความเห็นผู้เชี่ยวชาญ (เห็นด้วย / ควรแก้ / ไม่แน่ใจ)": 18, "ค่าที่เสนอแทน": 20,
    "เหตุผลหรือหมายเหตุของผู้เชี่ยวชาญ": 40, "ค่าคงที่ในโค้ด": 22, "ค่าในโค้ด": 22,
}
REVIEW_HEADER = "ความเห็นผู้เชี่ยวชาญ (เห็นด้วย / ควรแก้ / ไม่แน่ใจ)"
REVIEW = ("ความเห็นผู้เชี่ยวชาญ", "ค่าที่เสนอแทน", "เหตุผลหรือหมายเหตุ")
TYPE_FILL = {
    "มีแหล่งรองรับโดยตรง": "E2F0D9",
    "เลือกค่าในช่วงของแหล่งอ้างอิง": "FFF2CC",
    "กติกาโครงการ": "FCE4D6",
}


def main() -> None:
    rows = list(csv.DictReader(SOURCE.open(encoding="utf-8-sig")))
    columns = list(WIDTHS)
    wb = Workbook()
    ws = wb.active
    ws.title = "กติกาเครื่องคำนวณ"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws.append(columns)
    for cell in ws[1]:
        review = cell.value.startswith(REVIEW)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597" if review else "404040")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
    for row in rows:
        ws.append([row[c] for c in columns])
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = wrap
            cell.border = border
            if columns[cell.column - 1].startswith(REVIEW):
                cell.fill = PatternFill("solid", fgColor="DDEBF7")
        kind = r[columns.index("ประเภท")]
        kind.fill = PatternFill("solid", fgColor=TYPE_FILL.get(kind.value, "FFFFFF"))
    for i, name in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = WIDTHS[name]
    for name in ("ค่าคงที่ในโค้ด", "ค่าในโค้ด"):
        ws.column_dimensions[ws.cell(1, columns.index(name) + 1).column_letter].hidden = True
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    verdict = DataValidation(type="list", formula1='"เห็นด้วย,ควรแก้,ไม่แน่ใจ"', allow_blank=True)
    ws.add_data_validation(verdict)
    col = ws.cell(1, columns.index(REVIEW_HEADER) + 1).column_letter
    verdict.add(f"{col}2:{col}{len(rows) + 1}")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    sign = wb.create_sheet("ผู้ตรวจ")
    sign.column_dimensions["A"].width = 28
    sign.column_dimensions["B"].width = 60
    for label in ("ชื่อผู้ตรวจ", "ตำแหน่ง / วุฒิ / ใบอนุญาต", "หน่วยงาน", "วันที่ตรวจ",
                  "ความเห็นโดยรวม", "ลงชื่อ"):
        sign.append([label, ""])
    sign.append([])
    sign.append(["วิธีกรอก", ("เลือกความเห็นในคอลัมน์สีฟ้าของแต่ละข้อ ถ้าเลือก \"ควรแก้\" "
                 "กรุณาใส่ค่าที่เสนอและเหตุผล ข้อที่เป็น \"กติกาโครงการ\" คือค่าที่ผู้จัดทำเลือกเอง "
                 "ไม่มีงานวิจัยกำหนดค่านี้โดยตรง")])
    for r in sign.iter_rows():
        r[0].font = Font(bold=True)
        for cell in r:
            cell.alignment = wrap
    wb.save(TARGET)
    print(TARGET.relative_to(ROOT), len(rows), "rules")


if __name__ == "__main__":
    main()
